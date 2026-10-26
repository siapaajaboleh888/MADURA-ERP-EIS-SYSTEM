import sys
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN
from pptx.dml.color import RGBColor

def create_deck():
    prs = Presentation()
    # 16:9 Widescreen layout
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank_layout = prs.slide_layouts[6]

    # Color Palette
    C_DARK = RGBColor(15, 23, 42)      # Slate 900
    C_PRIMARY = RGBColor(245, 158, 11) # Amber 500
    C_ACCENT = RGBColor(56, 189, 248)  # Sky 400
    C_WHITE = RGBColor(255, 255, 255)
    C_MUTED = RGBColor(148, 163, 184)  # Slate 400
    C_CARD_BG = RGBColor(30, 41, 59)   # Slate 800

    def add_bg(slide, color):
        bg = slide.shapes.add_shape(1, 0, 0, prs.slide_width, prs.slide_height)
        bg.fill.solid()
        bg.fill.fore_color.rgb = color
        bg.line.fill.background()
        return bg

    # ==================== SLIDE 1: JUDUL ====================
    s1 = prs.slides.add_slide(blank_layout)
    add_bg(s1, C_DARK)

    # Title box
    tb = s1.shapes.add_textbox(Inches(1.5), Inches(2.2), Inches(10.3), Inches(3.0))
    tf = tb.text_frame
    tf.word_wrap = True

    p0 = tf.paragraphs[0]
    p0.text = "KENALAN DENGAN ERP & EIS"
    p0.font.size = Pt(44)
    p0.font.bold = True
    p0.font.color.rgb = C_PRIMARY
    p0.alignment = PP_ALIGN.CENTER

    p1 = tf.add_paragraph()
    p1.text = "Rahasia Cara Kerja Perusahaan Modern Tanpa Bikin Pusing!"
    p1.font.size = Pt(24)
    p1.font.color.rgb = C_WHITE
    p1.alignment = PP_ALIGN.CENTER

    p2 = tf.add_paragraph()
    p2.text = "\nStudi Kasus Santai: Bisnis Pabrik Garam Madura"
    p2.font.size = Pt(18)
    p2.font.color.rgb = C_ACCENT
    p2.alignment = PP_ALIGN.CENTER

    # ==================== SLIDE 2: MASALAH BISNIS TANPA SISTEM ====================
    s2 = prs.slides.add_slide(blank_layout)
    add_bg(s2, C_DARK)

    # Header
    tb = s2.shapes.add_textbox(Inches(1.0), Inches(0.8), Inches(11.3), Inches(1.2))
    tf = tb.text_frame
    p = tf.paragraphs[0]
    p.text = "Bayangkan Dulu: Pabrik Garam Tanpa Sistem Komputer"
    p.font.size = Pt(32)
    p.font.bold = True
    p.font.color.rgb = C_PRIMARY

    # 3 Cards
    cards = [
        ("Sales Terima Order", "Sales bilang ke pembeli: 'Bisa kirim 100 sak garam besok!' Padahal belum ngecek ke gudang sama sekali."),
        ("Gudang Panik", "Orang gudang kaget: 'Stok kita sisa 5 sak! Kenapa asal terima pesanan tanpa konfirmasi dulu?!'"),
        ("Keuangan Pusing", "Bagian keuangan bingung: 'Siapa aja yang udah bayar? Catatan di buku nota hilang ketumpahan kopi!'")
    ]
    for i, (title, desc) in enumerate(cards):
        left = Inches(1.0 + i * 3.9)
        box = s2.shapes.add_shape(1, left, Inches(2.3), Inches(3.6), Inches(4.2))
        box.fill.solid()
        box.fill.fore_color.rgb = C_CARD_BG
        box.line.color.rgb = RGBColor(71, 85, 105)

        tf_c = box.text_frame
        tf_c.word_wrap = True
        p_ct = tf_c.paragraphs[0]
        p_ct.text = f"\n❌ {title}"
        p_ct.font.size = Pt(20)
        p_ct.font.bold = True
        p_ct.font.color.rgb = RGBColor(251, 113, 133) # Rose

        p_cd = tf_c.add_paragraph()
        p_cd.text = f"\n{desc}"
        p_cd.font.size = Pt(15)
        p_cd.font.color.rgb = C_WHITE

    # ==================== SLIDE 3: APA ITU ERP? ====================
    s3 = prs.slides.add_slide(blank_layout)
    add_bg(s3, C_DARK)

    tb = s3.shapes.add_textbox(Inches(1.0), Inches(0.8), Inches(11.3), Inches(1.2))
    tf = tb.text_frame
    p = tf.paragraphs[0]
    p.text = "Solusi Masalah Ini Adalah: ERP!"
    p.font.size = Pt(32)
    p.font.bold = True
    p.font.color.rgb = C_PRIMARY

    # Main Card
    main_box = s3.shapes.add_shape(1, Inches(1.0), Inches(2.0), Inches(11.3), Inches(1.8))
    main_box.fill.solid()
    main_box.fill.fore_color.rgb = C_CARD_BG
    main_box.line.color.rgb = C_PRIMARY
    tf_m = main_box.text_frame
    tf_m.word_wrap = True
    p_m1 = tf_m.paragraphs[0]
    p_m1.text = "ERP = Enterprise Resource Planning"
    p_m1.font.size = Pt(22)
    p_m1.font.bold = True
    p_m1.font.color.rgb = C_PRIMARY

    p_m2 = tf_m.add_paragraph()
    p_m2.text = "Sederhananya: SATU aplikasi terpusat di mana SEMUA bagian (Sales, Gudang, Pabrik, Kasir) bekerja di tempat yang sama. Tidak ada lagi catat-catat manual yang terpisah!"
    p_m2.font.size = Pt(16)
    p_m2.font.color.rgb = C_WHITE

    # Flow Box
    erp_points = [
        ("1. Sales Klik Order", "Stok di gudang otomatis terkunci (di-booking) agar tidak dijual ke orang lain."),
        ("2. Gudang Tahu Otomatis", "Sistem memberi tahu staf gudang untuk menyiapkan kemasan garam."),
        ("3. Surat Jalan & Faktur Otomatis", "Saat truk jalan, faktur penagihan otomatis terbit ke bagian Finance.")
    ]
    for i, (title, desc) in enumerate(erp_points):
        left = Inches(1.0 + i * 3.9)
        box = s3.shapes.add_shape(1, left, Inches(4.2), Inches(3.6), Inches(2.5))
        box.fill.solid()
        box.fill.fore_color.rgb = C_CARD_BG
        box.line.color.rgb = C_ACCENT

        tf_c = box.text_frame
        tf_c.word_wrap = True
        p_t = tf_c.paragraphs[0]
        p_t.text = f"✨ {title}"
        p_t.font.size = Pt(17)
        p_t.font.bold = True
        p_t.font.color.rgb = C_ACCENT

        p_d = tf_c.add_paragraph()
        p_d.text = f"\n{desc}"
        p_d.font.size = Pt(14)
        p_d.font.color.rgb = C_WHITE

    # ==================== SLIDE 4: APA ITU EIS? ====================
    s4 = prs.slides.add_slide(blank_layout)
    add_bg(s4, C_DARK)

    tb = s4.shapes.add_textbox(Inches(1.0), Inches(0.8), Inches(11.3), Inches(1.2))
    tf = tb.text_frame
    p = tf.paragraphs[0]
    p.text = "Kalau ERP Sudah Ada, Kenapa Butuh EIS?"
    p.font.size = Pt(32)
    p.font.bold = True
    p.font.color.rgb = C_PRIMARY

    # Main Card
    main_box = s4.shapes.add_shape(1, Inches(1.0), Inches(2.0), Inches(11.3), Inches(1.8))
    main_box.fill.solid()
    main_box.fill.fore_color.rgb = C_CARD_BG
    main_box.line.color.rgb = C_ACCENT
    tf_m = main_box.text_frame
    tf_m.word_wrap = True
    p_m1 = tf_m.paragraphs[0]
    p_m1.text = "EIS = Executive Information System"
    p_m1.font.size = Pt(22)
    p_m1.font.bold = True
    p_m1.font.color.rgb = C_ACCENT

    p_m2 = tf_m.add_paragraph()
    p_m2.text = "Bos / Direktur TIDAK punya waktu memeriksa 1.000 nota belanja satu per satu!\nEIS adalah Dashboard pintar di HP/Laptop Bos yang merangkum jutaan data transaksi ERP menjadi angka dan grafik ringkas dalam 5 detik."
    p_m2.font.size = Pt(16)
    p_m2.font.color.rgb = C_WHITE

    # 4 KPI cards like EIS dashboard
    kpi_cards = [
        ("💰 Total Pendapatan", "Rp 2,45 Miliar", "Kas yang sudah lunas bulan ini"),
        ("⚠️ Piutang Tertunggak", "Rp 350 Juta", "Tagihan yang belum dibayar customer"),
        ("📦 Nilai Stok Gudang", "Rp 1,2 Miliar", "Total aset garam di 3 gudang"),
        ("🏆 Produk Terlaris", "Garam Halus 250g", "45.000 bungkus terjual")
    ]
    for i, (k_title, k_val, k_sub) in enumerate(kpi_cards):
        left = Inches(1.0 + i * 2.9)
        box = s4.shapes.add_shape(1, left, Inches(4.2), Inches(2.6), Inches(2.5))
        box.fill.solid()
        box.fill.fore_color.rgb = C_CARD_BG
        box.line.color.rgb = C_PRIMARY

        tf_c = box.text_frame
        tf_c.word_wrap = True
        p_t = tf_c.paragraphs[0]
        p_t.text = k_title
        p_t.font.size = Pt(14)
        p_t.font.color.rgb = C_MUTED

        p_v = tf_c.add_paragraph()
        p_v.text = f"\n{k_val}"
        p_v.font.size = Pt(20)
        p_v.font.bold = True
        p_v.font.color.rgb = C_PRIMARY

        p_s = tf_c.add_paragraph()
        p_s.text = f"\n{k_sub}"
        p_s.font.size = Pt(12)
        p_s.font.color.rgb = C_WHITE

    # ==================== SLIDE 5: PERBANDINGAN BERDAMPINGAN ====================
    s5 = prs.slides.add_slide(blank_layout)
    add_bg(s5, C_DARK)

    tb = s5.shapes.add_textbox(Inches(1.0), Inches(0.8), Inches(11.3), Inches(1.0))
    tf = tb.text_frame
    p = tf.paragraphs[0]
    p.text = "Perbedaan Utama: ERP vs EIS"
    p.font.size = Pt(32)
    p.font.bold = True
    p.font.color.rgb = C_PRIMARY

    # Side-by-side comparison boxes
    erp_box = s5.shapes.add_shape(1, Inches(1.0), Inches(1.9), Inches(5.4), Inches(4.8))
    erp_box.fill.solid()
    erp_box.fill.fore_color.rgb = C_CARD_BG
    erp_box.line.color.rgb = C_PRIMARY
    tf_e = erp_box.text_frame
    tf_e.word_wrap = True
    pe = tf_e.paragraphs[0]
    pe.text = "🏢 ERP (Operasional Bisnis)"
    pe.font.size = Pt(22)
    pe.font.bold = True
    pe.font.color.rgb = C_PRIMARY

    points_erp = [
        "\n• Pengguna: Staf harian (Sales, Admin Gudang, Mandor, Kasir)",
        "\n• Yang diinput: Nota, faktur, surat jalan, jumlah karung garam",
        "\n• Fokus: Kelancaran kerja harian antar bagian",
        "\n• Pertanyaan: 'Barang ini dikirim ke siapa dan kapan?'",
        "\n• Analogi: Mesin dan roda mobil yang memutar kendaraan"
    ]
    for pt in points_erp:
        p_item = tf_e.add_paragraph()
        p_item.text = pt
        p_item.font.size = Pt(14)
        p_item.font.color.rgb = C_WHITE

    eis_box = s5.shapes.add_shape(1, Inches(6.9), Inches(1.9), Inches(5.4), Inches(4.8))
    eis_box.fill.solid()
    eis_box.fill.fore_color.rgb = C_CARD_BG
    eis_box.line.color.rgb = C_ACCENT
    tf_i = eis_box.text_frame
    tf_i.word_wrap = True
    pi = tf_i.paragraphs[0]
    pi.text = "👑 EIS (Strategi Pimpinan)"
    pi.font.size = Pt(22)
    pi.font.bold = True
    pi.font.color.rgb = C_ACCENT

    points_eis = [
        "\n• Pengguna: Pimpinan puncak (CEO, Direktur, Pemilik Usaha)",
        "\n• Yang dilihat: Grafik omzet, untung/rugi, status piutang, KPI",
        "\n• Fokus: Pengambilan keputusan dan masa depan bisnis",
        "\n• Pertanyaan: 'Bulan ini untung berapa? Gudang mana yang lesu?'",
        "\n• Analogi: Speedometer dan GPS di depan mata sopir"
    ]
    for pt in points_eis:
        p_item = tf_i.add_paragraph()
        p_item.text = pt
        p_item.font.size = Pt(14)
        p_item.font.color.rgb = C_WHITE

    # ==================== SLIDE 6: KESIMPULAN / ANALOGI MOBIL ====================
    s6 = prs.slides.add_slide(blank_layout)
    add_bg(s6, C_DARK)

    tb = s6.shapes.add_textbox(Inches(1.0), Inches(0.8), Inches(11.3), Inches(1.0))
    tf = tb.text_frame
    p = tf.paragraphs[0]
    p.text = "Kesimpulan Super Gampang (Ingat Mobil Ini!)"
    p.font.size = Pt(32)
    p.font.bold = True
    p.font.color.rgb = C_PRIMARY

    # Big Banner
    box_car = s6.shapes.add_shape(1, Inches(1.0), Inches(2.0), Inches(11.3), Inches(3.8))
    box_car.fill.solid()
    box_car.fill.fore_color.rgb = C_CARD_BG
    box_car.line.color.rgb = RGBColor(16, 185, 129) # Emerald
    tf_car = box_car.text_frame
    tf_car.word_wrap = True

    pc0 = tf_car.paragraphs[0]
    pc0.text = "🚗 IBARAT SEBUAH MOBIL PERUSAHAAN:"
    pc0.font.size = Pt(24)
    pc0.font.bold = True
    pc0.font.color.rgb = RGBColor(52, 211, 153)

    pc1 = tf_car.add_paragraph()
    pc1.text = "\n1. ERP adalah Mesin, Roda, dan Transmisi"
    pc1.font.size = Pt(20)
    pc1.font.bold = True
    pc1.font.color.rgb = C_PRIMARY

    pc2 = tf_car.add_paragraph()
    pc2.text = "   Bekerja keras di bawah kap agar mobil bisa bergerak maju dan tidak mogok di jalan."
    pc2.font.size = Pt(16)
    pc2.font.color.rgb = C_WHITE

    pc3 = tf_car.add_paragraph()
    pc3.text = "\n2. EIS adalah Speedometer, Indikator Bensin, dan Layar Dashboard"
    pc3.font.size = Pt(20)
    pc3.font.bold = True
    pc3.font.color.rgb = C_ACCENT

    pc4 = tf_car.add_paragraph()
    pc4.text = "   Dilihat oleh pengemudi (Bos) agar tahu kecepatan mobil dan kapan harus isi bensin!"
    pc4.font.size = Pt(16)
    pc4.font.color.rgb = C_WHITE

    pc5 = tf_car.add_paragraph()
    pc5.text = "\n🤝 Keduanya saling membutuhkan: EIS tidak ada gunanya tanpa data dari ERP!"
    pc5.font.size = Pt(17)
    pc5.font.bold = True
    pc5.font.color.rgb = RGBColor(253, 224, 71)

    # Save
    out_path = "e:/TESTING/Materi_Mudah_ERP_dan_EIS.pptx"
    prs.save(out_path)
    print(f"Presentation saved to: {out_path}")

if __name__ == "__main__":
    create_deck()
