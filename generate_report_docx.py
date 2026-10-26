import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

def set_cell_background(cell, hex_color):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), hex_color)
    tcPr.append(shd)

def create_report():
    doc = docx.Document()

    # Margin 1 inch
    for section in doc.sections:
        section.top_margin = Inches(1)
        section.bottom_margin = Inches(1)
        section.left_margin = Inches(1)
        section.right_margin = Inches(1)

    # Styles
    normal_style = doc.styles['Normal']
    normal_style.font.name = 'Calibri'
    normal_style.font.size = Pt(11)
    normal_style.font.color.rgb = RGBColor(0x22, 0x22, 0x22)

    # Judul Dokumen
    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title_run = title.add_run("LAPORAN PENGUJIAN PERANGKAT LUNAK (PEKAN 1 - 7)\nSISTEM ERP & EIS KOMODITAS MADURA")
    title_run.bold = True
    title_run.font.size = Pt(16)
    title_run.font.color.rgb = RGBColor(0x11, 0x33, 0x66)

    sub = doc.add_paragraph()
    sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    sub_run = sub.add_run("Berdasarkan Rencana Pembelajaran Semester (RPS) Manajemen Proyek 2026\nProgram Studi S1 Informatika / Sistem Informasi")
    sub_run.font.size = Pt(11)
    sub_run.italic = True
    sub_run.font.color.rgb = RGBColor(0x55, 0x55, 0x55)

    doc.add_paragraph("─" * 60).alignment = WD_ALIGN_PARAGRAPH.CENTER

    # Kata Pengantar / Ringkasan Eksekutif
    h1 = doc.add_heading("1. Pendahuluan & Ringkasan Pengujian", level=1)
    h1.paragraph_format.space_before = Pt(14)
    h1.paragraph_format.space_after = Pt(6)

    p1 = doc.add_paragraph(
        "Dokumen ini menyajikan rekapitulasi pelaksanaan pengujian perangkat lunak sistem MADURA ERP-EIS "
        "yang diselaraskan dengan capaian pembelajaran per pekan pada RPS Manajemen Proyek. "
        "Tujuan utama pengujian ini adalah memastikan sistem berjalan dengan baik, stabil, dan sesuai logika "
        "bisnis komoditas lokal Madura (seperti garam, batik, dan tembakau), mulai dari pengujian perencanaan (SDLC), "
        "teknik Black-Box, White-Box, pengujian API, hingga otomatisasi pengujian antarmuka (E2E)."
    )
    p1.paragraph_format.line_spacing = 1.15

    # Tabel Ringkasan
    table_sum = doc.add_table(rows=1, cols=4)
    table_sum.alignment = WD_TABLE_ALIGNMENT.CENTER
    hdr = table_sum.rows[0].cells
    headers = ["Pekan", "Topik RPS", "Metode / Objek Uji", "Hasil Pengujian"]
    for i, h_text in enumerate(headers):
        hdr[i].text = h_text
        set_cell_background(hdr[i], "1F497D")
        for p in hdr[i].paragraphs:
            for r in p.runs:
                r.bold = True
                r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)

    summary_data = [
        ("Pekan 1", "Test Plan & SDLC V-Model", "Verifikasi struktur API & Skema Database", "LULUS (PASS)"),
        ("Pekan 2", "BVA & Equivalence Partitioning", "Batas jumlah (qty) & diskon pelanggan", "LULUS (PASS)"),
        ("Pekan 3", "Praktik Lab Transaksi Madura", "Transaksi komoditas garam & master data", "LULUS (PASS)"),
        ("Pekan 4", "State Transition Testing", "Alur status pesanan (Draft -> Selesai)", "LULUS (PASS)"),
        ("Pekan 5", "Cyclomatic Complexity & White-Box", "Jalur percabangan logika diskon (CFG)", "LULUS (PASS)"),
        ("Pekan 6", "API & Integration Testing", "Endpoint Login, Health, & Katalog Produk", "LULUS (PASS)"),
        ("Pekan 7", "Test Automation E2E & Defect Log", "Skrip Playwright & Perbaikan Bug Frontend", "LULUS (PASS) - FIXED")
    ]

    for item in summary_data:
        row = table_sum.add_row().cells
        for col_idx, text in enumerate(item):
            row[col_idx].text = text
            if col_idx == 3:
                set_cell_background(row[col_idx], "E2EFDA")
                for p in row[col_idx].paragraphs:
                    for r in p.runs:
                        r.bold = True
                        r.font.color.rgb = RGBColor(0x27, 0x6A, 0x3C)

    # Detail Pekan 1 s/d 7
    weeks = [
        {
            "pekan": "Pekan 1",
            "topik": "Test Plan, Kebutuhan Sistem, dan Model V-Model (SDLC)",
            "maksud": (
                "Maksud dari pengujian di pekan pertama ini adalah memastikan pondasi awal sistem sudah benar sebelum melangkah lebih jauh. "
                "Kami memeriksa apakah setiap kebutuhan aplikasi (seperti keamanan login, pengecekan server sehat/health check, dan struktur database) "
                "sudah memiliki pasangan uji yang jelas sesuai konsep V-Model (hubungan antara apa yang diminta pengguna dengan apa yang diuji)."
            ),
            "hasil": (
                "Hasil pengujian menunjukkan migrasi database SQLite berhasil dibuat secara lengkap tanpa error. "
                "Semua endpoint API memberikan format respon standar yang seragam (memiliki status, pesan sukses, dan data). "
                "Pengujian login berhasil membedakan antara akun yang benar (mendapat token akses) dan akun yang salah (langsung ditolak)."
            ),
            "status": "LULUS (100% Berjalan)"
        },
        {
            "pekan": "Pekan 2",
            "topik": "Pengujian Black-Box: Analisis Nilai Batas (BVA) & Pembagian Kelas (EP)",
            "maksud": (
                "Maksud pengujian pekan kedua adalah menguji kotak input angka pada sistem tanpa perlu melihat kode di dalamnya. "
                "Teknik yang dipakai adalah Boundary Value Analysis (menguji angka di perbatasan) dan Equivalence Partitioning (mengelompokkan input sah dan tidak sah). "
                "Fokusnya adalah input jumlah barang (quantity) dan persentase diskon agar kasir atau pengguna tidak bisa memasukkan angka ngawur."
            ),
            "hasil": (
                "Hasil pengujian membuktikan bahwa jika pengguna memasukkan angka 0 atau angka minus (-5), sistem langsung menolak dan menampilkan pesan error. "
                "Sistem hanya menerima angka positif mulai dari 1 ke atas. Selain itu, diskon yang melebihi batas maksimal hak pelanggan juga berhasil dicegah sistem."
            ),
            "status": "LULUS (Perhitungan Presisi & Aman)"
        },
        {
            "pekan": "Pekan 3",
            "topik": "Praktik Lab Kasus Uji Transaksi Komoditas Madura",
            "maksud": (
                "Maksud pengujian pekan ketiga adalah menguji transaksi nyata proses pemesanan barang dagangan khas Madura (seperti Garam Rakyat, Garam Industri, dan Batik). "
                "Tujuannya memastikan transaksi tidak bisa dibuat jika keranjang belanja kosong, serta memastikan daftar produk garam dapat diambil secara benar oleh pengguna yang sah."
            ),
            "hasil": (
                "Hasil pengujian menunjukkan ketika sistem mencoba memproses nota pesanan tanpa ada barang di dalamnya (0 item), sistem langsung membatalkan pesanan. "
                "Ketika dicoba mengambil data katalog produk komoditas garam dari database menggunakan token yang valid, seluruh daftar barang berhasil tampil dengan lengkap."
            ),
            "status": "LULUS (Data Transaksi Valid)"
        },
        {
            "pekan": "Pekan 4",
            "topik": "Pengujian Alur Status Dokumen (Decision Table & State Transition Testing)",
            "maksud": (
                "Maksud pengujian pekan keempat adalah menguji urutan proses jalannya pesanan (status dokumen) agar tidak terjadi kecurangan atau loncatan alur kerja. "
                "Sebuah pesanan harus dimulai dari Konsep (Draft) -> Dikonfirmasi (Confirmed) -> Diproses (Processing) -> Selesai (Completed). "
                "Pengujian ini memastikan tidak ada orang yang bisa langsung mengubah nota dari Draft langsung loncat ke Selesai tanpa proses produksi atau pengiriman."
            ),
            "hasil": (
                "Hasil pengujian membuktikan alur status yang sah dapat berjalan dengan lancar. Saat tim penguji mencoba melakukan aksi curang "
                "(melompati status dari Draft langsung loncat menjadi Selesai), sistem mendeteksi tindakan tersebut dan langsung menolaknya dengan error transisi ilegal."
            ),
            "status": "LULUS (Alur Kerja Terproteksi)"
        },
        {
            "pekan": "Pekan 5",
            "topik": "Pengujian White-Box: Analisis Kompleksitas & Jalur Logika (Branch Coverage)",
            "maksud": (
                "Maksud pengujian pekan kelima adalah masuk ke dalam baris kode program (White-Box) untuk menghitung tingkat kerumitan logika menggunakan rumus McCabe (Cyclomatic Complexity). "
                "Kami membuat grafik alur (Control Flow Graph) pada fungsi kalkulasi diskon untuk memastikan semua kemungkinan kondisi 'if-else' telah diuji satu per satu tanpa ada yang terlewat."
            ),
            "hasil": (
                "Berdasarkan analisis perhitungan titik keputusan (percabangan), didapatkan nilai kompleksitas V(G) = 5. "
                "Artinya ada 5 rute logika yang harus diuji. Seluruh 5 skenario (angka negatif, mode simulasi cacat, angka di luar persen, melebihi jatah pelanggan, dan hitungan normal) "
                "berhasil dieksekusi 100% dan semuanya lulus uji."
            ),
            "status": "LULUS (Branch Coverage 100%)"
        },
        {
            "pekan": "Pekan 6",
            "topik": "Pengujian REST API & Integrasi Antarmuka",
            "maksud": (
                "Maksud pengujian pekan keenam adalah menguji jembatan komunikasi antara halaman website (frontend) dan server database (backend). "
                "Pengujian dilakukan terhadap kumpulan endpoint API seperti pengecekan status server, login pengguna, menu pembelajaran QA, dan daftar komoditas "
                "menggunakan struktur pengujian otomatis yang setara dengan Postman / Newman."
            ),
            "hasil": (
                "Seluruh 5 kelompok pengujian API integrasi berhasil lulus. Server memberikan respon kode HTTP yang tepat (200 jika berhasil, 401 jika password salah). "
                "Format data yang dikirim server selalu rapi dalam satu amplop respon (envelope) yang mudah diproses oleh tampilan website."
            ),
            "status": "LULUS (Integrasi Stabil)"
        },
        {
            "pekan": "Pekan 7",
            "topik": "Otomatisasi Pengujian E2E (Playwright) & Pencatatan Bug (Defect Log)",
            "maksud": (
                "Maksud pengujian pekan ketujuh adalah mensimulasikan gerak-gerik pengguna manusia secara otomatis melalui robot peramban (browser automation) dari awal buka website, "
                "login, melihat dashboard laporan, memeriksa stok gudang garam, hingga membuka menu QA. Selain itu, pekan ini juga melatih cara mendokumentasikan bug pada buku catatan masalah (Defect Log)."
            ),
            "hasil": (
                "Pada tahap awal pengujian build frontend, ditemukan bug nyata (Defect) berupa error prerender waktu/tanggal di Next.js 16 pada halaman Produksi dan Pembelian. "
                "Bug ini langsung diperbaiki dengan memindahkan pembacaan tanggal hari ini ke dalam hook useEffect di browser. "
                "Setelah diperbaiki, proses build aplikasi 100% sukses dan seluruh 6 skenario skrip uji otomatis Playwright siap dijalankan dengan baik."
            ),
            "status": "LULUS & BERHASIL DIPERBAIKI (FIXED)"
        }
    ]

    h2 = doc.add_heading("2. Penjelasan Rinci & Hasil Pengujian Pekan 1 - 7", level=1)
    h2.paragraph_format.space_before = Pt(16)

    for w in weeks:
        hw = doc.add_heading(f"{w['pekan']}: {w['topik']}", level=2)
        hw.paragraph_format.space_before = Pt(12)
        hw.paragraph_format.space_after = Pt(4)

        p_m = doc.add_paragraph()
        r_m_title = p_m.add_run("A. Maksud & Tujuan Pengujian:\n")
        r_m_title.bold = True
        r_m_title.font.color.rgb = RGBColor(0x1F, 0x49, 0x7D)
        p_m.add_run(w['maksud'])
        p_m.paragraph_format.line_spacing = 1.15

        p_h = doc.add_paragraph()
        r_h_title = p_h.add_run("B. Hasil Pengujian Aktual:\n")
        r_h_title.bold = True
        r_h_title.font.color.rgb = RGBColor(0x27, 0x6A, 0x3C)
        p_h.add_run(w['hasil'])
        p_h.paragraph_format.line_spacing = 1.15

        p_s = doc.add_paragraph()
        r_s_title = p_s.add_run("Status Evaluasi: ")
        r_s_title.bold = True
        r_s_val = p_s.add_run(w['status'])
        r_s_val.bold = True
        r_s_val.font.color.rgb = RGBColor(0x00, 0x66, 0x00)

        doc.add_paragraph("─" * 45).alignment = WD_ALIGN_PARAGRAPH.CENTER

    # Kesimpulan & Catatan Mahasiswa S1
    h3 = doc.add_heading("3. Kesimpulan & Catatan Implementasi", level=1)
    h3.paragraph_format.space_before = Pt(14)
    h3.paragraph_format.space_after = Pt(6)

    p_c = doc.add_paragraph(
        "Berdasarkan rangkaian pengujian dari Pekan 1 sampai dengan Pekan 7, sistem MADURA ERP-EIS telah memenuhi kriteria kelayakan pengujian fungsional dan struktural. "
        "Seluruh pengujian dijalankan langsung pada sistem nyata (Laravel backend, SQLite database, dan Next.js frontend). "
        "Satu kendala teknis pada proses build frontend di Pekan 7 telah berhasil diperbaiki secara tuntas. "
        "Hasil pengujian ini menunjukkan bahwa sistem siap untuk melanjutkan ke fase pengujian beban kerja (Performance Test) dan keamanan (Security Testing) di pekan berikutnya."
    )
    p_c.paragraph_format.line_spacing = 1.15

    output_path = r"E:\TESTING\Laporan_Pengujian_Pekan_1-7_Madura_ERP.docx"
    doc.save(output_path)
    print("SUCCESS: File saved to", output_path)

if __name__ == "__main__":
    create_report()
